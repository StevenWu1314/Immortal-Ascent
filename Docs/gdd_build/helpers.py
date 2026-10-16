"""Shared python-docx helpers for the Immortal Ascent GDD build."""
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

_doc = None
_widths_cache = {}
_list_counter = 0


def _reset_list():
    """Any non-list paragraph ends the current numbered list.

    Word's built-in List Number style shares one numbering definition for the whole
    document, which made the count run on across sections. Numbers are written as
    literal text instead, and reset whenever the list is interrupted.
    """
    global _list_counter
    _list_counter = 0


def _repeat_header(row):
    """Mark a table row as a header row so it repeats on every page it spans."""
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _apply_widths(table, widths):
    """Set per-column widths in the table grid as well as on every cell.

    LibreOffice (and Word in fixed-layout mode) sizes columns from w:tblGrid, so
    setting cell widths alone is not enough.
    """
    table.autofit = False
    tbl = table._tbl
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        cols = grid.findall(qn("w:gridCol"))
        for col, width in zip(cols, widths):
            col.set(qn("w:w"), str(int(width * 1440)))
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 1440)))
    for row in table.rows:
        for i, width in enumerate(widths):
            row.cells[i].width = Inches(width)


def init(doc):
    global _doc
    _doc = doc
    return doc


def _clean(text):
    return " ".join(str(text).split())


def title(text, sub=None):
    _reset_list()
    p = _doc.add_heading(text, level=0)
    if sub:
        q = _doc.add_paragraph()
        r = q.add_run(_clean(sub))
        r.italic = True
        r.font.size = Pt(11)
    return p


def h1(text):
    _reset_list()
    return _doc.add_heading(text, level=1)


def h2(text):
    _reset_list()
    return _doc.add_heading(text, level=2)


def h3(text):
    _reset_list()
    return _doc.add_heading(text, level=3)


def p(text=""):
    _reset_list()
    return _doc.add_paragraph(_clean(text))


def tag(label, text):
    """Paragraph with a bold lead-in label, e.g. 'Decision: ...'."""
    _reset_list()
    para = _doc.add_paragraph()
    run = para.add_run(label + " ")
    run.bold = True
    para.add_run(_clean(text))
    return para


def b(text):
    _reset_list()
    return _doc.add_paragraph(_clean(text), style="List Bullet")


def n(text):
    """Numbered item. Numbering restarts whenever the list is interrupted."""
    global _list_counter
    _list_counter += 1
    para = _doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.32)
    para.paragraph_format.first_line_indent = Inches(-0.25)
    para.add_run("%d. %s" % (_list_counter, _clean(text)))
    return para


def note(text):
    _reset_list()
    para = _doc.add_paragraph()
    run = para.add_run(_clean(text))
    run.italic = True
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    return para


def t(headers, rows, widths=None):
    """Table Grid table. widths = list of inches."""
    _reset_list()
    table = _doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, head in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(_clean(head))
        run.bold = True
        run.font.size = Pt(9.5)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(_clean(val))
            run.font.size = Pt(9.5)
    if widths:
        _apply_widths(table, widths)
    _repeat_header(table.rows[0])
    _doc.add_paragraph()
    return table


def kv(rows):
    """Two-column label/value table without a header row."""
    _reset_list()
    table = _doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        run = cells[0].paragraphs[0].add_run(_clean(label))
        run.bold = True
        run.font.size = Pt(9.5)
        run2 = cells[1].paragraphs[0].add_run(_clean(value))
        run2.font.size = Pt(9.5)
        cells[0].width = Inches(1.9)
        cells[1].width = Inches(4.6)
    _doc.add_paragraph()
    return table


def page_break():
    _doc.add_page_break()
