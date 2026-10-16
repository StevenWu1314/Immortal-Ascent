"""Build Immortal-Ascent-GDD-v2.docx from the content modules."""
import os
import sys
import datetime

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

import helpers
import part1
import part2
import part3
import part4
import to_markdown

OUT = os.path.abspath(os.path.join(HERE, "..", "Immortal-Ascent-GDD-v2.docx"))
MD_OUT = os.path.abspath(os.path.join(HERE, "..", "Immortal-Ascent-GDD-v2.md"))


def setup(doc):
    for s in doc.sections:
        s.left_margin = Inches(0.85)
        s.right_margin = Inches(0.85)
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11.5)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0x1F, 0x33, 0x46)
        st.font.bold = True
    title_style = doc.styles["Title"]
    title_style.font.name = "Calibri"
    title_style.font.size = Pt(26)
    title_style.font.color.rgb = RGBColor(0x1F, 0x33, 0x46)


def cover(doc):
    today = datetime.date.today().isoformat()
    helpers.title("Immortal Ascent",
                  "Game Design Document v2 - Rogue Asian Fantasy")
    sub = doc.add_paragraph()
    run = sub.add_run("Draft for team review")
    run.bold = True
    run.font.size = Pt(12)
    helpers.p("This document is a full replacement for the original mind-map GDD "
              "(Immortal Ascent Overhaul.pdf). It contains the design as decided in a structured "
              "interview with the designer, written out in enough detail to implement from, plus "
              "two appendices covering what is still undecided and what the current build does not "
              "yet do.")
    helpers.kv([
        ("Version", "2.0 (draft)"),
        ("Date", today),
        ("Prepared for", "Steven (design + code), Dewei (art), Jo (UI)"),
        ("Source", "Immortal Ascent Overhaul.pdf, Docs/ItemSystem.md, Docs/InscriptionSystem.md, "
                   "Docs/Refactoring-Plan.md, and a read-only survey of Assets/Scripts"),
        ("Project", "Rogue Asian Fantasy, Unity 2022.3.62f2"),
        ("Status", "Decisions locked; values marked 'proposed' need playtesting"),
    ])
    helpers.page_break()


def main():
    doc = Document()
    setup(doc)
    helpers.init(doc)
    cover(doc)
    for module in (part1, part2, part3, part4):
        module.build()
    doc.save(OUT)
    print("wrote", OUT)
    print("wrote", to_markdown.export(OUT, MD_OUT))


if __name__ == "__main__":
    main()
